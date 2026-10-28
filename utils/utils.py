from collections import defaultdict
from datetime import datetime
import re
from tkinter import filedialog
from openpyxl import Workbook, load_workbook
import sys
from pathlib import Path

from openpyxl.styles import PatternFill
from utils.sheetStyles import *
from openpyxl.utils import get_column_letter
from database.data import vaccineList


def getDniFromExcel(file_route):
    wb = load_workbook(file_route, read_only=True, data_only=True)
    ws = wb.active

    dnis = []
    dnis_set = set()
    errors = []

    for index_row, (value,) in enumerate(
        ws.iter_rows(min_col=1, max_col=1, values_only=True),
        start=1,
    ):
        if value is None:
            continue  # Ignora filas vacías

        dni_text = str(value).strip()

        # Si Excel guardó el DNI como número, recupera ceros iniciales.
        if isinstance(value, int):
            dni_text = f"{value:08d}"

        if re.fullmatch(r"\d{8}", dni_text):
           if dni_text not in dnis_set:
                dnis.append(dni_text)
                dnis_set.add(dni_text)
        else:
            errors.append(
                f"Fila {index_row}: debe contener un DNI de 8 dígitos."
            )

    wb.close()
    return dnis, errors


def export_excel(data, cpmsList, text):
    if not data:
        return None
    
    vaccineGroupByDniAndLab = defaultdict(
        lambda: defaultdict(
            lambda:defaultdict( list))
    )
    patients = {}
    filename = text[10:]
    for dni, patient, fn, cpms, lab, date in data:
        if dni not in patients:
            patients[dni] = {
                "paciente": patient,
                "fn": fn,
            }

        # Una atención sin LAB también es válida.
        lab_key = str(lab).strip() if lab is not None else ""
        if not lab_key:
            lab_key = "SIN_LAB"

        if date:
            vaccineGroupByDniAndLab[dni][cpms][lab_key].append(date)
    # Identificar los LAB únicos, incluyendo SIN_LAB.
    labsByCpms = {
        cpms: sorted({
            lab
            for group in vaccineGroupByDniAndLab.values()
            for lab in group.get(cpms, {})
        })
        for cpms in cpmsList    
    }
    columnsByCpmsAndLab = {}
     # Determinar cuántas columnas necesita cada LAB.
    for cpms in cpmsList:
        for lab in labsByCpms[cpms]:
            columnsByCpmsAndLab[(cpms, lab)] = max(
                (
                    len(group.get(cpms, {}).get(lab, []))
                    for group in vaccineGroupByDniAndLab.values()
                ),
                default=0,
            )
    
    output_route = filedialog.asksaveasfilename(
        title="Guardar resultados",
        defaultextension=".xlsx",
        filetypes=[("Archivo Excel", "*.xlsx")],
        initialfile=f"resultado_{'-'.join(cpmsList)}_{filename}"
    )

    if not output_route:
        return None

    book = Workbook()
    sheet = book.active
    sheet.title = "Atenciones"
    
    # Crear cabeceras dinámicas.
    sheet.append(["", "", "", ""])
    sheet.append(["DNI", "Paciente", "FN", "Edad"])
    col = 5
    for index,cpms in enumerate(cpmsList):
        vaccine_name = vaccineList.get(cpms, cpms)
        firstCol = col
        color = vaccineColors[index % len(vaccineColors)]
        
        for lab in labsByCpms[cpms]:
            columnLabSize = columnsByCpmsAndLab[(cpms, lab)]
            
            for i in range(columnLabSize):
                sufix = "" if i == 0 else f"({i})"
                sheet.cell(
                    row=2,
                    column=col,
                    value=f"{lab}{sufix}"
                )
            col += 1
        endCol = col-1
        
        if(endCol >= firstCol):
            sheet.cell(
                row=1,
                column=firstCol,
                value=vaccine_name
            )
            
            if(endCol > firstCol):
                sheet.merge_cells(
                    start_row = 1,
                    start_column=firstCol,
                    end_row=1,
                    end_column=endCol
                )
                
            for colIndex in range(firstCol, endCol + 1):
                for row in (1, 2):
                    current_cell = sheet.cell(row=row, column=colIndex)
                    current_cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor=color
                    )
                    current_cell.font = headerStyle
                    current_cell.alignment = alignment
                    current_cell.border = border
    
    for colIndex in range(1, 5):
        current_cell = sheet.cell(row=2, column=colIndex)
        current_cell.fill = PatternFill(
            fill_type="solid",
            fgColor="334155"
        )
        current_cell.font = headerStyle
        current_cell.alignment = alignment
        current_cell.border = border

    # Altura de los encabezados.
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 23
    
    for dni, datos in patients.items():
        row = [
            dni, 
            datos["paciente"] or "", 
            datos["fn"] or "", 
            None
            ]

        for cpms in cpmsList:
            for lab in labsByCpms[cpms]:
                dates = sorted(
                    vaccineGroupByDniAndLab[dni]
                    .get(cpms, {})
                    .get(lab, []),
                    key=sort_dates
                )

                amount = columnsByCpmsAndLab[(cpms, lab)]

                for number in range(amount):
                    row.append(
                        dates[number]
                        if number < len(dates)
                        else ""
                    )

        sheet.append(row)
        current_row = sheet.max_row
        sheet.cell(
            row=current_row,
            column=4,
            value=f'=IF(C{current_row}="","",DATEDIF(C{current_row},TODAY(),"Y"))'
        )


    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = (
        f"A2:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    )
    
    sheet.column_dimensions["A"].width = 13
    sheet.column_dimensions["B"].width = 34
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 10

    for colIndex in range(5, sheet.max_column + 1):
        sheet.column_dimensions[
            get_column_letter(colIndex)
        ].width = 15
        
    book.save(output_route)

    return output_route

def sort_dates(fecha):
    if not fecha:
        return datetime.min

    texto = str(fecha).strip()

    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%Y%m%d"):
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue

    return datetime.min

def resource_path(relative_path):
    if hasattr(sys, "_MEIPASS"):
        base_path = Path(sys._MEIPASS)
    else:
        base_path = Path(__file__).resolve().parent.parent

    return base_path / relative_path